"""Create a separate Word copy with editable Office Math and explicit bilingual fonts."""
from pathlib import Path
import re, json, zipfile
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('.')
SOURCE = ROOT / '预测模型与变量说明.md'
TARGET = ROOT / '预测模型与变量说明.docx'
QA = Path('output\\text_cleanup\\20261014T203329199980+0800_all_docs\\word_qa')
QA.mkdir(parents=True, exist_ok=True)
XSL = E.XSLT(E.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
MML = 'http://www.w3.org/1998/Math/MathML'
NS = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

# Each mapping preserves the source equation; long equations are split at additive terms.
FORMULAS = {
'T̂_t = T_(t−1)^2 / T_(t−2)': r'\hat{T}_t=\frac{T_{t-1}^{2}}{T_{t-2}}',
'T̂_formal1 = T0': r'\hat{T}_{\text{formal1}}=T_0',
'T0 × (T_practice1/T0)': r'T_0\times\frac{T_{\text{practice1}}}{T_0}',
'T̂_(t+1) = T_t × (1+g_T)': r'\hat{T}_{t+1}=T_t(1+g_T)',
'g_T = T̂_(t+1)/T_t−1': r'g_T=\frac{\hat{T}_{t+1}}{T_t}-1',
'u_i = −e_c ln(P_i/P0_i) + b_G ln(1+G_i/1000) + b_H ln(1+H_i/1000) + b_R ln(f_R(R_i+ρ_R R_(i,t−1)))': [
r'u_i=-e_c\ln\frac{P_i}{P_{0,i}}+b_G\ln(1+\frac{G_i}{1000})',
r'\quad+b_H\ln(1+\frac{H_i}{1000})',
r'\quad+b_R\ln[f_R(R_i+\rho_R R_{i,t-1})]'],
'ŝ_i = exp(u_i)/sum_(j∈J) exp(u_j)': r'\hat{s}_i=\frac{\exp(u_i)}{\sum_{j\in J}\exp(u_j)}',
'D̂_i = T̂ × ŝ_i': r'\hat{D}_i=\hat{T}\hat{s}_i',
'ln(D_i)−组内均值[ln(D)] = (x_i−组内均值[x])·β': r'\ln(D_i)-\overline{\ln(D)}=(x_i-\bar{x})\cdot\beta',
'D̂_i = exp(α_c + u_i)': r'\hat{D}_i=\exp(\alpha_c+u_i)',
'ŝ_i = 1/n_J': r'\hat{s}_i=\frac{1}{n_J}',
'D̂_i = T̂/n_J': r'\hat{D}_i=\frac{\hat{T}}{n_J}',
'D̂_i,oracle = T_actual × ŝ_i': r'\hat{D}_{i,\text{oracle}}=T_{\text{actual}}\hat{s}_i',
'D̂_i = exp[α_c − e_c ln(P_i/P0_c) + b_G ln(1+G_i/1000) + b_H ln(1+H_i/1000) + b_R ln(1+(R_i+ρ_R R_(i,t−1))/1000)]': [
r'\ln\hat{D}_i=\alpha_c-e_c\ln\frac{P_i}{P_{0,c}}+b_G\ln(1+\frac{G_i}{1000})',
r'\quad+b_H\ln(1+\frac{H_i}{1000})',
r'\quad+b_R\ln(1+\frac{R_i+\rho_R R_{i,t-1}}{1000})'],
'x_R = R_t + ρ_R R_(t−1)': r'x_R=R_t+\rho_R R_{t-1}',
'f_R(x_R) = a_R + b_curve,R (x_R+c_R)^d_R': r'f_R(x_R)=a_R+b_{\text{curve},R}(x_R+c_R)^{d_R}',
'x_O = O_t + ρ_O O_(t−1)': r'x_O=O_t+\rho_O O_{t-1}',
'f_O(x_O) = a_O + b_curve,O x_O/(c_O+x_O)': r'f_O(x_O)=a_O+b_{\text{curve},O}\frac{x_O}{c_O+x_O}',
'η̂ = θ × f_O(x_O)': r'\hat{\eta}=\theta f_O(x_O)',
'C_t = 0.85 C_(t−1) + I_t': r'C_t=0.85C_{t-1}+I_t',
'K_machine = p_I I_t': r'K_{\text{machine}}=p_I I_t',
'q̂ = min(C_t a/100, sum_c D̂_c)': r'\hat{q}=\min(\frac{C_t a}{100},\sum_c\hat{D}_c)',
'a_candidate = 100 min(sum_c D̂_c/C_t,1)': r'a_{\text{candidate}}=100\min(\frac{\sum_c\hat{D}_c}{C_t},1)',
'Ŝ_c = min(D̂_c,Q_remaining)': r'\hat{S}_c=\min(\hat{D}_c,Q_{\text{remaining}})',
'q1=min(q̂,C/3)': r'q_1=\min(\hat{q},\frac{C}{3})',
'q2=min(max(q̂−C/3,0),C/3)': r'q_2=\min(\max(\hat{q}-\frac{C}{3},0),\frac{C}{3})',
'q3=max(q̂−2C/3,0)': r'q_3=\max(\hat{q}-\frac{2C}{3},0)',
'N̂ = q̂/η̂': r'\hat{N}=\frac{\hat{q}}{\hat{\eta}}',
'W = (q1+κ2 q2+κ3 q3)w/η̂': r'W=\frac{(q_1+\kappa_2 q_2+\kappa_3 q_3)w}{\hat{\eta}}',
'F = max((1−δ)N_(t−1)−N̂,0)wσ': r'F=\max((1-\delta)N_{t-1}-\hat{N},0)w\sigma',
'V=sum_c(P_c Ŝ_c)': r'V=\sum_c P_c\hat{S}_c',
'B=q̂(v+Δv)': r'B=\hat{q}(v+\Delta v)',
'Z=sum_c(k_c Ŝ_c)': r'Z=\sum_c k_c\hat{S}_c',
'A_marketing=G+H+50': r'A_{\text{marketing}}=G+H+50',
'A_overhead=4000+F+R+O': r'A_{\text{overhead}}=4000+F+R+O',
'EBITDA=V−B−W−Z−A_marketing−A_overhead': r'\text{EBITDA}=V-B-W-Z-A_{\text{marketing}}-A_{\text{overhead}}',
'Dep_machine=0.15(Book_machine,0+K_machine)': r'\text{Dep}_{\text{machine}}=0.15(\text{Book}_{\text{machine},0}+K_{\text{machine}})',
'Dep_building=0.07 Book_building,0': r'\text{Dep}_{\text{building}}=0.07\text{Book}_{\text{building},0}',
'EBIT=EBITDA−Dep_machine−Dep_building': r'\text{EBIT}=\text{EBITDA}-\text{Dep}_{\text{machine}}-\text{Dep}_{\text{building}}',
'r_L = r_base + a_r exp(−b_r E/A)': r'r_L=r_{\text{base}}+a_r\exp(-b_r\frac{E}{A})',
'r_L = r_base + a_r ((L+L_short)/E)^b_r': r'r_L=r_{\text{base}}+a_r(\frac{L+L_{\text{short}}}{E})^{b_r}',
'L=L0+ΔL': r'L=L_0+\Delta L',
'J_L=(L0+L)r_L/2': r'J_L=\frac{L_0+L}{2}r_L',
'AR=V×30/365；AP=B×30/365': r'\text{AR}=V\frac{30}{365}\quad;\quad\text{AP}=B\frac{30}{365}',
'Cash_pre=Cash0+EBITDA−J_L+(AR0−AR)+(AP−AP0)−K_machine−K_building+ΔL': [
r'\text{Cash}_{\text{pre}}=\text{Cash}_0+\text{EBITDA}-J_L+(\text{AR}_0-\text{AR})',
r'\quad+(\text{AP}-\text{AP}_0)-K_{\text{machine}}-K_{\text{building}}+\Delta L'],
'J_C=max((Cash0+Cash)/2,0)r_C': r'J_C=\max(\frac{\text{Cash}_0+\text{Cash}}{2},0)r_C',
'PBT=EBIT+J_C−J_L': r'\text{PBT}=\text{EBIT}+J_C-J_L',
'Tax=max(PBT−Loss0,0)τ': r'\text{Tax}=\max(\text{PBT}-\text{Loss}_0,0)\tau',
'Profit=PBT−Tax': r'\text{Profit}=\text{PBT}-\text{Tax}',
'Cash=Cash_pre+J_C−Tax': r'\text{Cash}=\text{Cash}_{\text{pre}}+J_C-\text{Tax}',
'E=E0+Profit': r'E=E_0+\text{Profit}',
'Book_machine=Book_machine,0+K_machine−Dep_machine': r'\text{Book}_{\text{machine}}=\text{Book}_{\text{machine},0}+K_{\text{machine}}-\text{Dep}_{\text{machine}}',
'Book_building=Book_building,0+K_building−Dep_building': r'\text{Book}_{\text{building}}=\text{Book}_{\text{building},0}+K_{\text{building}}-\text{Dep}_{\text{building}}',
'A=Book_machine+Book_building+AR+Cash': r'A=\text{Book}_{\text{machine}}+\text{Book}_{\text{building}}+\text{AR}+\text{Cash}',
'A=L+AP+E': r'A=L+\text{AP}+E',
'Loss_end=max(Loss0−max(PBT,0),0)+max(−PBT,0)': r'\text{Loss}_{\text{end}}=\max(\text{Loss}_0-\max(\text{PBT},0),0)+\max(-\text{PBT},0)',
'L_short,needed=max(Cash_min−Cash,0)': r'L_{\text{short},\text{needed}}=\max(\text{Cash}_{\text{min}}-\text{Cash},0)',
'g_D = D̂_(t+1)/D̂_t−1': r'g_D=\frac{\hat{D}_{t+1}}{\hat{D}_t}-1',
}

def ml(tag, text=None, *children):
    n=E.Element('{'+MML+'}'+tag)
    if text is not None: n.text=text
    n.extend(children)
    return n

class Tex:
    """Small strict parser for this document's fixed equation notation."""
    symbols={'rho':'ρ','alpha':'α','eta':'η','theta':'θ','Delta':'Δ','kappa':'κ',
             'delta':'δ','sigma':'σ','beta':'β','tau':'τ','pi':'π','in':'∈','times':'×','cdot':'·'}
    def __init__(self,text):
        self.tokens=re.findall(r'\\[A-Za-z]+|\d+(?:\.\d+)?|[A-Za-z]+|[^\s]',text)
        self.i=0
    def group(self):
        if self.tokens[self.i]=='{':
            self.i+=1; out=self.expr('}'); self.i+=1; return out
        return self.atom()
    def atom(self):
        t=self.tokens[self.i]; self.i+=1
        if t=='{':
            out=self.expr('}'); self.i+=1
        elif t==r'\frac': out=ml('mfrac',None,self.group(),self.group())
        elif t in (r'\hat',r'\bar',r'\overline'):
            out=ml('mover',None,self.group(),ml('mo','^' if t==r'\hat' else '¯')); out.set('accent','true')
        elif t==r'\text':
            assert self.tokens[self.i]=='{'; self.i+=1; words=[]
            while self.tokens[self.i]!='}': words.append(self.tokens[self.i]); self.i+=1
            self.i+=1; out=ml('mi',''.join(words)); out.set('mathvariant','normal')
        elif t==r'\sum': out=ml('mo','∑'); out.set('largeop','true')
        elif t in (r'\ln',r'\exp',r'\min',r'\max'):
            out=ml('mi',t[1:]); out.set('mathvariant','normal')
        elif t==r'\quad': out=ml('mspace'); out.set('width','1em')
        elif t.startswith('\\'):
            assert t[1:] in self.symbols,t
            v=self.symbols[t[1:]]; out=ml('mo' if v in '∈×·' else 'mi',v)
        elif re.fullmatch(r'\d+(?:\.\d+)?',t): out=ml('mn',t)
        else: out=ml('mi' if t.isalpha() else 'mo','−' if t=='-' else t)
        sub=sup=None
        while self.i<len(self.tokens) and self.tokens[self.i] in ['_','^']:
            which=self.tokens[self.i]; self.i+=1; value=self.group()
            if which=='_': sub=value
            else: sup=value
        if sub is not None and sup is not None: out=ml('msubsup',None,out,sub,sup)
        elif sub is not None: out=ml('msub',None,out,sub)
        elif sup is not None: out=ml('msup',None,out,sup)
        return out
    def expr(self,end=None):
        nodes=[]
        while self.i<len(self.tokens) and self.tokens[self.i]!=end: nodes.append(self.atom())
        return ml('mrow',None,*nodes)

def fonts(rpr):
    f=rpr.find(qn('w:rFonts'))
    if f is None: f=OxmlElement('w:rFonts'); rpr.insert(0,f)
    for attr in ['ascii','hAnsi','cs']: f.set(qn('w:'+attr),'Times New Roman')
    f.set(qn('w:eastAsia'),'宋体')
    for key in list(f.attrib):
        if 'Theme' in key: del f.attrib[key]

def math(latex):
    parser=Tex(latex); m=ml('math',None,parser.expr()); assert parser.i==len(parser.tokens)
    out=XSL(m).getroot()
    for run in out.xpath('.//m:r',namespaces=NS):
        props=run.find(qn('w:rPr'))
        if props is None: props=OxmlElement('w:rPr'); run.insert(0,props)
        fonts(props)
    return out

def plain_latex(text):
    text=text.replace('−','-')
    text=re.sub(r'([A-Za-zα-ωΑ-Ω])̂',r'\\hat{\1}',text)
    text=re.sub(r'P0_([ic])',r'P_{0,\1}',text)
    text=re.sub(r'([A-Za-zα-ωΑ-Ω]+)_([A-Za-z0-9]+(?:,[A-Za-z0-9]+)?)',r'\1_{\2}',text)
    return text

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
for name,size in [('Normal',11),('Title',20),('Heading 1',14),('Heading 2',12)]:
    style=doc.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0); fonts(style.element.get_or_add_rPr())
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.18
    if name.startswith('Heading'): style.paragraph_format.space_before=Pt(12)
    for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.widow_control=True
mathpr=OxmlElement('m:mathPr'); font=OxmlElement('m:mathFont'); font.set(qn('m:val'),'Times New Roman'); mathpr.append(font)
old=doc.settings.element.find(qn('m:mathPr'))
if old is not None: doc.settings.element.remove(old)
doc.settings.element.append(mathpr)
doc.core_properties.title='预测模型与变量说明'
doc.core_properties.subject='Cesim预测模型公式 变量 单位和适用条件'
doc.core_properties.author=''

equations=[]
def add_math(p,latex): p._p.append(math(latex)); equations.append(latex)
def equation(source,after=None):
    assert source in FORMULAS, source
    values=FORMULAS[source]
    if isinstance(values,str): values=[values]
    for k,value in enumerate(values):
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.keep_with_next=k<len(values)-1 or after is not None
        add_math(p,value)
    if after: rich(doc.add_paragraph(),after)

def rich(p,text):
    text=text.replace('**','')
    for part in re.split(r'(`[^`]+`|\[[^\]]+\]\([^)]+\))',text):
        if part.startswith('`'):
            formula=part[1:-1]
            if formula not in FORMULAS:
                run=p.add_run(formula); fonts(run._r.get_or_add_rPr()); continue
            values=FORMULAS[formula]; assert isinstance(values,str)
            add_math(p,values)
        elif part.startswith('[') and (link:=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)):
            h=OxmlElement('w:hyperlink')
            target=link.group(2)
            if target.startswith('C:/'): target='file:///'+target
            elif not target.startswith(('http:', 'https:')): target=SOURCE.parent.joinpath(target).as_uri()
            h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r'); props=OxmlElement('w:rPr'); fonts(props)
            r.append(props); t=OxmlElement('w:t'); t.text=link.group(1); r.append(t); h.append(r); p._p.append(h)
        elif part:
            # Format identifiers in explanations consistently with their formula symbols.
            pattern=r'([A-Za-zα-ωΑ-Ω]+̂?(?:_[A-Za-zα-ωΑ-Ω0-9]+(?:,[0-9]+)?|_\(t\+1\))|[DSqη]̂|\b(?:P0|Cash0|AR0|AP0|Loss0|L0|E0|q[123]|κ[23])\b)'
            for piece in re.split(pattern,part):
                if piece and re.fullmatch(pattern,piece):
                    piece=piece.replace('_(t+1)','_{t+1}')
                    piece=re.sub(r'^(P|Cash|AR|AP|Loss|L|E)(0)$',r'\1_{\2}',piece)
                    piece=re.sub(r'^(q|κ)([123])$',r'\1_{\2}',piece)
                    add_math(p,plain_latex(piece))
                elif piece:
                    run=p.add_run(piece); fonts(run._r.get_or_add_rPr())

lines=SOURCE.read_text(encoding='utf-8').splitlines(); pos=0; section=''
while pos<len(lines):
    line=lines[pos].strip(); pos+=1
    if not line: continue
    if line.startswith('**主题导航：**'): continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:],style='Title'); continue
    if line.startswith('## '):
        section=line[3:]
        heading=re.sub(r'[.（）/、：]+',' ',section).strip()
        doc.add_paragraph(heading,style='Heading 1'); continue
    if line.startswith('|'):
        table_lines=[line]
        while pos<len(lines) and lines[pos].strip().startswith('|'):
            table_lines.append(lines[pos].strip()); pos+=1
        rows=[[c.strip() for c in x.strip('|').split('|')] for x in table_lines]
        rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
        if section.startswith(('10.','12.')):
            for row in rows[1:]: equation(row[0],row[1])
            continue
        table=doc.add_table(rows=0,cols=len(rows[0])); table.style='Table Grid'; table.autofit=False
        widths=[4.0,13.0] if len(rows[0])==2 else [3.2,8.7,5.1]
        if section=='来源与参数入口': widths=[7.5,9.5]
        for row_no,values in enumerate(rows):
            cells=table.add_row().cells
            trpr=cells[0]._tc.getparent().get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if row_no==0: trpr.append(OxmlElement('w:tblHeader'))
            for col,(cell,value) in enumerate(zip(cells,values)):
                cell.width=Cm(widths[col]); p=cell.paragraphs[0]
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.08
                if row_no==0: p.paragraph_format.keep_with_next=True
                if col==0 and row_no>0 and re.fullmatch(r'[A-Za-z0-9α-ωΑ-Ω_、/，,∈̂−. ]+',value):
                    add_math(p,plain_latex(value))
                else: rich(p,value)
                for run in p.runs:
                    run.font.size=Pt(10)
                    if row_no==0: run.bold=True
                if row_no==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F0F0F0'); cell._tc.get_or_add_tcPr().append(shade)
        if section!='来源与参数入口': doc.add_paragraph().paragraph_format.space_after=Pt(0)
        continue
    # Standalone long equations are displayed as structured formulas.
    code=re.findall(r'`([^`]+)`',line)
    if code and any(isinstance(FORMULAS.get(c),list) for c in code):
        assert len(code)==1
        equation(code[0]); continue
    rich(doc.add_paragraph(),line)

# Footer page number is unobtrusive and uses a native Word field.
p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run(); fonts(r._r.get_or_add_rPr()); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
doc.save(TARGET)

# Trust-boundary checks: native math exists, structures are preserved, font settings explicit.
with zipfile.ZipFile(TARGET) as z:
    xml=E.fromstring(z.read('word/document.xml'))
    assert len(xml.xpath('//m:oMath',namespaces=NS))==len(equations)>60
    assert len(xml.xpath('//m:f',namespaces=NS))>15
    assert len(xml.xpath('//m:sSub',namespaces=NS))>30
    assert not xml.xpath('//w:drawing',namespaces=NS)
    assert not xml.xpath('//w:t[contains(text(),"`")]',namespaces=NS)
    for props in xml.xpath('//m:r/w:rPr/w:rFonts',namespaces=NS):
        assert props.get(qn('w:ascii'))=='Times New Roman' and props.get(qn('w:eastAsia'))=='宋体'
    result={'source':SOURCE.name,'output':TARGET.name,'native_equations':len(equations),
            'fractions':len(xml.xpath('//m:f',namespaces=NS)),
            'subscripts':len(xml.xpath('//m:sSub',namespaces=NS)),
            'font_latin':'Times New Roman','font_chinese':'宋体','formula_images':0,
            'equation_tex_for_audit':equations}
    (QA/'structure_check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k!='equation_tex_for_audit'},ensure_ascii=False))
